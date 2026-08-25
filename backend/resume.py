"""简历解析：PDF / DOCX / TXT / MD -> 纯文本。"""
from __future__ import annotations

import io
import re

MAX_CHARS = 20000  # 超长简历截断，避免撑爆上下文


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.strip()
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n（简历过长，已截断）"
    return text


def extract(filename: str, data: bytes) -> str:
    name = (filename or "").lower()

    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
        text = "\n".join(pages)
        if not text.strip():
            raise ValueError("这个 PDF 提取不到文字（可能是扫描件/图片版），请导出文字版 PDF 或改传 Word/Markdown")
        return _clean(text)

    if name.endswith(".docx"):
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:  # 简历常用表格排版
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return _clean("\n".join(parts))

    if name.endswith(".doc"):
        raise ValueError("不支持旧版 .doc，请另存为 .docx 或导出 PDF")

    if name.endswith((".txt", ".md", ".markdown")):
        for enc in ("utf-8", "gbk", "utf-16"):
            try:
                return _clean(data.decode(enc))
            except UnicodeDecodeError:
                continue
        raise ValueError("文本编码无法识别")

    raise ValueError("仅支持 PDF / DOCX / TXT / MD 格式")


def summarize_for_ui(text: str, limit: int = 180) -> str:
    """返回给前端展示的开头片段。"""
    one = re.sub(r"\s+", " ", text).strip()
    return one[:limit] + ("…" if len(one) > limit else "")
